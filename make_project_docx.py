# -*- coding: utf-8 -*-
"""生成给老师看的简明汇报稿（大白话版）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_run_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.get_or_add_rPr()
    rfonts = run._element.rPr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run(text), "黑体", 20, True)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run(text), "黑体", 14, True)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(0.74)
    set_run_font(p.add_run(text), "宋体", 12)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(text), "宋体", 12)


doc = Document()

# 页面与正文默认样式
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

add_title(doc, "项目目前遇到的实质性问题")

add_body(
    doc,
    "这个项目做的是高铁站外交通疏导。真正难的地方不在写程序，而在于现场很多数据没有，"
    "每个站的情况又不一样。下面按几个方面简单说。",
)

add_heading(doc, "一、路段监控")
add_bullet(doc, "现在说的“每段路都有监控”，其实是每段配一路演示摄像头，用一两段样片视频循环播放，真摄像头还没有接。")
add_bullet(doc, "哪路视频对应哪段路，必须一条条登记清楚；登记错了，数据就会算到别的路上。")
add_bullet(doc, "从外面进站的车流可以拿视频数据来校正；站内路段不能灌，不然就等于“放录像”，推荐路线就不起作用了。")

add_heading(doc, "二、落客点")
add_bullet(doc, "肇庆东站的落客区全在地下一层，没有高架落客平台，做模拟要单独按地下车库来建。")
add_bullet(doc, "落客点、上客点的位置没有官方坐标，现在图上是按道路走向估出来的“演示点”，不是现场量的。")
add_bullet(doc, "车停着不动要分三种：排队算堵、正常落客不算堵、违停算事件，不能混在一起。")

add_heading(doc, "三、停车场")
add_bullet(doc, "肇庆东的停车数据是几年前的了：地下 533 个私家车位加 80 个出租车蓄车位，露天 660 个小车位。"
                "露天场 2025 年 11 月开始改扩建，改完有多少车位还没公布。")
add_bullet(doc, "出租车只有 80 个蓄车位，高峰期很容易满，这正是“预约放车”要解决的地方。")
add_bullet(doc, "广州南站停车场编号比较乱，12 个场子却有 14 个名字，P9 的用途也变过。")
add_bullet(doc, "广州南站出租车是两级：蓄车场约 770 个位，上客区同时只能接 40 到 100 台；上客区一满，车就堆在蓄车场。")

add_heading(doc, "四、道路有几条车道")
add_bullet(doc, "61 段路里，只有 10 段知道有几条车道，另外 51 段没有数据，只能空着，不能编。")
add_bullet(doc, "双向的路也没写两边各几条，只能合着算。")
add_bullet(doc, "现场核对车道数的表还没交上来，所以现在只能标“推测”或“默认值”。")

add_heading(doc, "五、双向道路")
add_bullet(doc, "61 段里 33 段是双向、28 段是单向。很多路没标单行，就默认当双向，还得补 33 条反方向的路，不然车开进去出不来。")
add_bullet(doc, "双向路两个方向共用一个通行能力，预约时分不清进站和离站各占多少。")
add_bullet(doc, "掉头只在断头路允许；左转限制没有数据，现在只记录、不拿它算路线。")
add_bullet(doc, "全部进出组合 67 对，能走通 48 对；走不通的 19 对是单行和方向造成的，属正常，但都如实记下来了。")

add_heading(doc, "六、其他")
add_bullet(doc, "地图坐标有三套（GPS、高德、百度），不转换的话，位置和路能差几十上百米。")
add_bullet(doc, "红绿灯只核实了一个路口，其余只能推测。")
add_bullet(doc, "客流只有一天的总数，没有分时段数据，也不知道私家车、出租车、网约车各占多少。")
add_bullet(doc, "站前还在修珠肇高铁，围挡会改道；晚上 9 点后公交只剩一条线，黑车、乱停跟这个有关。")

add_heading(doc, "总结")
add_body(
    doc,
    "这些问题不是写代码能解决的。现在的做法是：缺数据的地方先用演示数据顶上，"
    "同时在文档和界面上标明“这是演示、缺什么、等谁来补”，没有把没做的说成做了。",
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.space_before = Pt(18)
set_run_font(p.add_run("2026 年 10 月"), "宋体", 12)

doc.save("项目目前遇到的实质性问题.docx")
print("saved")
